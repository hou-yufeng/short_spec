from __future__ import annotations

import argparse
import importlib
import os
import sys
from pathlib import Path

from html_shortspec_format import normalize_short_spec, validate_workbook
from html_source_rules import feature_values, finalize_short_spec

VENDORED_DEPS = Path(__file__).resolve().parent / "_third_batch_pydeps"
if VENDORED_DEPS.exists():
    sys.path.insert(0, str(VENDORED_DEPS))


def collect_html_paths(source_dir: str | None, pattern: str, html_files: list[str] | None) -> list[Path]:
    paths = [Path(path).resolve() for path in html_files] if html_files else sorted(Path(source_dir or "").glob(pattern))
    if not paths:
        raise SystemExit("No HTML source files found.")
    return paths


def write_xlsx(path: Path, rows: list[list[str]], config: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All Products"
    for index, row in enumerate(rows):
        if index:
            raw_ports = (config in {"com", "con", "smb"} and row[2] == "Ports") or (config == "dt" and "Ports" in row[2])
            row = [*row[:3], normalize_short_spec(row[3] if raw_ports else finalize_short_spec(row[3]))]
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for column, width in {"A": 36, "B": 22, "C": 28, "D": 72}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    work_dir = path.parent / "_work"
    work_dir.mkdir(parents=True, exist_ok=True)
    temporary = work_dir / f"{path.stem}.format_check_{os.getpid()}.xlsx"
    workbook.save(temporary)
    workbook.close()
    issues = validate_workbook(temporary, require_l2_order=True)
    if issues:
        for issue in issues:
            print(f"FORMAT_ISSUE\t{issue.sheet}\t{issue.row}\t{issue.code}\t{issue.message}", file=sys.stderr)
        raise RuntimeError(f"Final ShortSpec format check found {len(issues)} issue(s): {temporary}")
    os.replace(temporary, path)
    print("FORMAT_ISSUES\t0")


PRODUCT_RULE_PACKAGES = {
    "com": "commercial_laptop", "con": "consumer_laptop", "smb": "smb_laptop",
    "tab": "tablet", "dt": "desktop", "ts": "thinkstation",
}


def product_name(path: Path) -> str:
    return path.stem.removesuffix("_Spec")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate third-batch HTML ShortSpec features.")
    parser.add_argument("--config", required=True, choices=sorted(PRODUCT_RULE_PACKAGES))
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--source-dir")
    source.add_argument("--html-files", nargs="+")
    parser.add_argument("--glob", default="*.html")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--output-xlsx", required=True)
    args = parser.parse_args()
    module = importlib.import_module(f"html_product_rules.{PRODUCT_RULE_PACKAGES[args.config]}.third_batch")
    paths = collect_html_paths(args.source_dir, args.glob, args.html_files)
    if args.limit is not None:
        paths = paths[:args.limit]
    rows = [["Product", "L1 Feature", "L2 Feature", "Short Spec"]]
    for path in paths:
        html_text = path.read_text(encoding="utf-8-sig", errors="replace")
        product_rows = [
            [product_name(path), l1, l2, short_spec]
            for l1, l2, short_spec in module.generate(feature_values(html_text), html_text)
        ]
        rows.extend(sorted(product_rows, key=lambda row: row[2].casefold()))
    output = Path(args.output_xlsx)
    output.parent.mkdir(parents=True, exist_ok=True)
    write_xlsx(output, rows, args.config)
    print(f"PRODUCTS\t{len(paths)}")
    print(f"ROWS\t{len(rows)-1}")
    print(f"WORKBOOK\t{output}")


if __name__ == "__main__":
    main()
