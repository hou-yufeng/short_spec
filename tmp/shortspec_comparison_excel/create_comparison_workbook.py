from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).with_name("pydeps")))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT_PATH = Path(r"D:\shortspec_generator\Why_Short_Specifications_Matter.xlsx")
SHEET_NAME = "Why Short Specifications Matter"

ROWS = [
    ["Dimension", "Full Product Specification", "Short Specification", "Business Value"],
    [
        "Primary purpose",
        "Provides complete technical and configuration details",
        "Highlights key configurations and selling points",
        "Matches information depth to the audience's needs",
    ],
    [
        "Typical audience",
        "Product experts, engineers, and technical reviewers",
        "Sales, marketing, channels, customers, and management",
        "Enables faster cross-functional communication",
    ],
    [
        "Content volume",
        "Comprehensive and often lengthy",
        "Concise and focused",
        "Reduces reading and search time",
    ],
    [
        "Information structure",
        "Includes all options, conditions, notes, and technical details",
        "Consolidates the most relevant information using business rules",
        "Makes product differentiation easier to identify",
    ],
    [
        "Ease of use",
        "Suitable for detailed research and technical validation",
        "Suitable for quick review, comparison, and communication",
        "Speeds up product launch and sales enablement activities",
    ],
    [
        "Consistency requirement",
        "Accuracy of complete source data",
        "Accuracy, clarity, and consistent wording across products",
        "Strengthens product-information quality and brand consistency",
    ],
    [
        "Update impact",
        "Changes may be distributed across many detailed sections",
        "Requires controlled summarization after source changes",
        "Makes traceable, rule-based maintenance essential",
    ],
]


def apply_border(cell, *, top=False, bottom=False, left=False, right=False) -> None:
    color = "D9E2F3"
    thin = Side(style="thin", color=color)
    cell.border = Border(
        top=thin if top else Side(style=None),
        bottom=thin if bottom else Side(style=None),
        left=thin if left else Side(style=None),
        right=thin if right else Side(style=None),
    )


def create_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Why Short Specifications Matter"
    sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="0B1F3A")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:D2")
    sheet["A2"] = "Short specifications transform detailed product data into concise, business-ready information."
    sheet["A2"].font = Font(name="Aptos", size=11, italic=True, color="1F3A5F")
    sheet["A2"].fill = PatternFill("solid", fgColor="EAF1F8")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 36

    for row_index, values in enumerate(ROWS, start=4):
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            apply_border(
                cell,
                top=True,
                bottom=True,
                left=True,
                right=True,
            )
            if row_index == 4:
                cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif col_index == 1:
                cell.font = Font(name="Aptos", size=10, bold=True, color="1F3A5F")
                cell.fill = PatternFill("solid", fgColor="F2F6FA")
            elif col_index == 4:
                cell.font = Font(name="Aptos", size=10, color="1F5136")
                cell.fill = PatternFill("solid", fgColor="EEF6F2")

    sheet.row_dimensions[4].height = 30
    for row_index in range(5, 12):
        sheet.row_dimensions[row_index].height = 52

    widths = {"A": 22, "B": 37, "C": 37, "D": 40}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells("A13:D13")
    sheet["A13"] = (
        "Key message: Short specifications are not simplified copies of full specifications. "
        "They are structured business outputs that transform detailed product data into concise, "
        "consistent, and usable information."
    )
    sheet["A13"].font = Font(name="Aptos", size=10, bold=True, color="704D00")
    sheet["A13"].fill = PatternFill("solid", fgColor="FFF4D6")
    sheet["A13"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A13"].border = Border(
        top=Side(style="thin", color="E6C76B"),
        bottom=Side(style="thin", color="E6C76B"),
        left=Side(style="thin", color="E6C76B"),
        right=Side(style="thin", color="E6C76B"),
    )
    sheet.row_dimensions[13].height = 42

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_title_rows = "4:4"
    sheet.sheet_properties.outlinePr.summaryBelow = False

    workbook.save(OUTPUT_PATH)


def verify_workbook() -> None:
    workbook = load_workbook(OUTPUT_PATH, data_only=False)
    sheet = workbook[SHEET_NAME]
    assert sheet["A1"].value == "Why Short Specifications Matter"
    assert sheet["A4"].value == "Dimension"
    assert sheet["D11"].value == "Makes traceable, rule-based maintenance essential"
    assert sheet["A13"].value.startswith("Key message:")
    assert sheet.max_row == 13 and sheet.max_column == 4
    assert sheet.freeze_panes == "A5"
    assert sheet["A1"].fill.fgColor.rgb == "000B1F3A"
    assert sheet["A4"].font.bold is True
    print(f"Verified workbook: {OUTPUT_PATH}")
    print(f"Sheet: {sheet.title}; populated range: A1:D13")


if __name__ == "__main__":
    create_workbook()
    verify_workbook()
