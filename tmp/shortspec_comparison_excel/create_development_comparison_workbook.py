from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).with_name("pydeps")))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUTPUT_PATH = Path(r"D:\shortspec_generator\Traditional_vs_AI_Enabled_Development.xlsx")
SHEET_NAME = "Traditional vs AI Development"

ROWS = [
    ["Dimension", "Traditional Development", "AI-Enabled Development"],
    [
        "Capability dependency",
        "Relies heavily on specialist development teams.",
        "Enables business experts and developers to collaborate more directly.",
    ],
    [
        "Requirement translation",
        "Long handoff path from business requirements to implementation.",
        "Business language can be converted into rules, examples, and implementation drafts faster.",
    ],
    [
        "Prototyping speed",
        "Requires a full development cycle before validation.",
        "Allows rapid validation of parsing logic, rules, and output formats.",
    ],
    [
        "Rule maintenance",
        "Rules may be implicit in code or individual experience.",
        "Rules, samples, and edge cases can be made explicit, reusable, and easier to review.",
    ],
    [
        "Testing",
        "Test coverage depends heavily on development capacity and manual test design.",
        "AI can assist with edge-case design, regression samples, and output comparisons.",
    ],
    [
        "Knowledge retention",
        "Knowledge often depends on individual experience and documentation quality.",
        "Expert knowledge can be captured as executable rules and test assets.",
    ],
    [
        "Team collaboration",
        "Business requests, development implements, and testing validates in separate stages.",
        "Business, engineering, testing, and AI form a faster, iterative feedback loop.",
    ],
    [
        "Change response",
        "Changes typically require analysis, development scheduling, implementation, and retesting.",
        "Changes can be scoped, prototyped, tested, and refined more quickly with human review.",
    ],
]


def all_borders(color: str) -> Border:
    thin = Side(style="thin", color=color)
    return Border(top=thin, bottom=thin, left=thin, right=thin)


def create_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5"

    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Traditional Development vs. AI-Enabled Development"
    sheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="0B1F3A")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet.merge_cells("A2:C2")
    sheet["A2"] = "AI changes how capabilities are built and improved while human experts retain ownership of quality and business rules."
    sheet["A2"].font = Font(name="Aptos", size=11, italic=True, color="1F3A5F")
    sheet["A2"].fill = PatternFill("solid", fgColor="EAF1F8")
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 36

    for row_index, values in enumerate(ROWS, start=4):
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            cell.border = all_borders("D9E2F3")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Aptos", size=10, color="1F2937")
            if row_index == 4:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif col_index == 1:
                cell.fill = PatternFill("solid", fgColor="F2F6FA")
                cell.font = Font(name="Aptos", size=10, bold=True, color="1F3A5F")
            elif col_index == 2:
                cell.fill = PatternFill("solid", fgColor="FAFAFA")
            else:
                cell.fill = PatternFill("solid", fgColor="EEF6F2")
                cell.font = Font(name="Aptos", size=10, color="1F5136")

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 54
    sheet.column_dimensions["C"].width = 58
    sheet.row_dimensions[4].height = 30
    for row_index in range(5, 13):
        sheet.row_dimensions[row_index].height = 52

    sheet.merge_cells("A14:C14")
    sheet["A14"] = "Key message: AI does not replace professional teams. It helps teams turn expert knowledge into faster, more repeatable, and more scalable delivery capability."
    sheet["A14"].font = Font(name="Aptos", size=10, bold=True, color="704D00")
    sheet["A14"].fill = PatternFill("solid", fgColor="FFF4D6")
    sheet["A14"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A14"].border = all_borders("E6C76B")
    sheet.row_dimensions[14].height = 42

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_title_rows = "4:4"
    workbook.save(OUTPUT_PATH)


def verify_workbook() -> None:
    workbook = load_workbook(OUTPUT_PATH, data_only=False)
    sheet = workbook[SHEET_NAME]
    assert sheet["A1"].value == "Traditional Development vs. AI-Enabled Development"
    assert sheet["A4"].value == "Dimension"
    assert sheet["C12"].value.startswith("Changes can be scoped")
    assert sheet["A14"].value.startswith("Key message:")
    assert sheet.max_row == 14 and sheet.max_column == 3
    assert sheet.freeze_panes == "A5"
    assert sheet["A4"].font.bold is True
    assert sheet["C5"].fill.fgColor.rgb == "00EEF6F2"
    print(f"Verified workbook: {OUTPUT_PATH}")
    print(f"Sheet: {sheet.title}; populated range: A1:C14")


if __name__ == "__main__":
    create_workbook()
    verify_workbook()
