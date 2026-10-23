"""Final text cleanup and workbook checks for HTML-source ShortSpec output."""

from __future__ import annotations

import argparse
import html
import re
import sys
from dataclasses import dataclass
from pathlib import Path


HEADERS = ("Product", "L1 Feature", "L2 Feature", "Short Spec")
TRADEMARK_MARKS = "\u2122\u00ae"
HTML_ENTITY = re.compile(r"&(?:[a-z][a-z0-9]+|#\d+|#x[0-9a-f]+);", flags=re.I)


@dataclass(frozen=True)
class FormatIssue:
    sheet: str
    row: int
    code: str
    message: str


def normalize_short_spec(value: str) -> str:
    """Apply the shared final output rule without changing feature-specific text."""
    return html.unescape(value).translate(str.maketrans("", "", TRADEMARK_MARKS))


def validate_workbook(path: Path, require_l2_order: bool = False) -> list[FormatIssue]:
    """Check final Excel cells after all product-line rules have been applied."""
    vendored = Path(__file__).resolve().parent / "_third_batch_pydeps"
    if vendored.exists() and str(vendored) not in sys.path:
        sys.path.insert(0, str(vendored))
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    issues: list[FormatIssue] = []
    try:
        for sheet in workbook.worksheets:
            header = tuple(sheet.cell(1, column).value for column in range(1, 5))
            if header != HEADERS:
                issues.append(FormatIssue(sheet.title, 1, "header", f"Expected columns: {', '.join(HEADERS)}"))
                continue
            previous_l2_by_product: dict[str, str] = {}
            for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=4, values_only=True), start=2):
                if not any(value is not None and str(value).strip() for value in cells):
                    continue
                for column, value in enumerate(cells):
                    if value is None or not str(value).strip():
                        issues.append(FormatIssue(sheet.title, row_number, "empty_cell", f"{HEADERS[column]} is empty"))
                product, l2_feature = cells[0], cells[2]
                if require_l2_order and isinstance(product, str) and isinstance(l2_feature, str):
                    key = l2_feature.casefold()
                    previous = previous_l2_by_product.get(product)
                    if previous is not None and key < previous:
                        issues.append(FormatIssue(sheet.title, row_number, "l2_order", "L2 Feature is not in A-Z order for this product"))
                    previous_l2_by_product[product] = key
                value = cells[3]
                if not isinstance(value, str):
                    continue
                if any(mark in value for mark in TRADEMARK_MARKS):
                    issues.append(FormatIssue(sheet.title, row_number, "trademark", "Short Spec contains \u2122 or \u00ae"))
                if HTML_ENTITY.search(value):
                    issues.append(FormatIssue(sheet.title, row_number, "html_entity", "Short Spec contains an undecoded HTML entity"))
                if value != value.strip() or "\r" in value or "\t" in value:
                    issues.append(FormatIssue(sheet.title, row_number, "whitespace", "Short Spec has outer whitespace, a tab, or CR"))
    finally:
        workbook.close()
    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="Check an HTML-source ShortSpec workbook's final format.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--require-l2-order", action="store_true", help="Check A-Z L2 order within each product.")
    args = parser.parse_args()
    issues = validate_workbook(args.workbook, require_l2_order=args.require_l2_order)
    for issue in issues:
        print(f"{issue.sheet}\t{issue.row}\t{issue.code}\t{issue.message}")
    print(f"FORMAT_ISSUES\t{len(issues)}")
    return 1 if issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
