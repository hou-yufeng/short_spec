from __future__ import annotations

import argparse
import importlib
import os
import sys
from pathlib import Path

from html_source_rules import feature_values

VENDORED_DEPS = Path(__file__).resolve().parent / "_third_batch_pydeps"
if VENDORED_DEPS.exists():
    sys.path.insert(0, str(VENDORED_DEPS))


def collect_html_paths(source_dir: str | None, pattern: str, html_files: list[str] | None) -> list[Path]:
    paths = [Path(path).resolve() for path in html_files] if html_files else sorted(Path(source_dir or "").glob(pattern))
    if not paths:
        raise SystemExit("No HTML source files found.")
    return paths


def write_xlsx(path: Path, rows: list[list[str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All Products"
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for column, width in {"A": 36, "B": 22, "C": 28, "D": 72}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


PRODUCT_RULE_PACKAGES = {
    "com": "commercial_laptop", "con": "consumer_laptop", "smb": "smb_laptop",
    "tab": "tablet", "dt": "desktop", "ts": "thinkstation",
}


def product_name(path: Path) -> str:
    return path.stem.removesuffix("_Spec")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate third-batch HTML ShortSpec features.")
    parser.add_argument("--config", required=True, choices=sorted(PRODUCT_RULE_PACKAGES))
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--source-dir")
    source.add_argument("--html-files", nargs="+")
    parser.add_argument("--glob", default="*.html")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--output-xlsx", required=True)
    args = parser.parse_args()
    module = importlib.import_module(f"html_product_rules.{PRODUCT_RULE_PACKAGES[args.config]}.third_batch")
    paths = collect_html_paths(args.source_dir, args.glob, args.html_files)
    if args.limit is not None:
        paths = paths[:args.limit]
    rows = [["Product", "L1 Feature", "L2 Feature", "Short Spec"]]
    for path in paths:
        html_text = path.read_text(encoding="utf-8-sig", errors="replace")
        for l1, l2, short_spec in module.generate(feature_values(html_text)):
            rows.append([product_name(path), l1, l2, short_spec])
    output = Path(args.output_xlsx)
    output.parent.mkdir(parents=True, exist_ok=True)
    write_xlsx(output, rows)
    print(f"PRODUCTS\t{len(paths)}")
    print(f"ROWS\t{len(rows)-1}")
    print(f"WORKBOOK\t{output}")


if __name__ == "__main__":
    main()
